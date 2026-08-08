"""Fill the workbook's Actual Result + Verdict columns from the observed run.

Only records what was observed (Playwright UI run + backend harness). Does not
touch application code. Also updates the Bugs sheet Status where this run
changed our confidence, and adds a Run Info sheet.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FILE = "pinelabs_bug_report.xlsx"
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
PASS = PatternFill("solid", fgColor="A9D08E")
CONFIRM = PatternFill("solid", fgColor="F4B183")   # bug reproduced
REFUTE = PatternFill("solid", fgColor="D9D9D9")
INFO = PatternFill("solid", fgColor="BDD7EE")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)

# tc -> (actual, verdict, fill)
ACTUAL = {
    "TC-01": (
        "Opened the submitted unpaid invoice in Desk. Both 'Pay on Machine' and "
        "'Send Payment Link' buttons rendered in the toolbar.",
        "PASS — positive control works", PASS),
    "TC-02": (
        "Opened the Paid invoice (outstanding 0). Neither button rendered.",
        "PASS — buttons correctly hidden when outstanding=0", PASS),
    "TC-03": (
        "Opened the DRAFT invoice (status pill 'Draft', docstatus 0). BOTH buttons "
        "rendered. outstanding_amount is populated to 1000 on the saved draft, so the "
        "JS amount gate (>0) passes.",
        "BUG-13 REFUTED — buttons DO appear on a saved draft; the 'drafts have "
        "outstanding_amount=0' premise is false in this ERPNext version", REFUTE),
    "TC-04": (
        "Opened Pinelabs Transaction PLT-2026-00002 (PENDING). Changed Status to "
        "FAILED via the Select and saved (Ctrl+S). NO error was shown; the save was "
        "accepted and the row persisted as status=FAILED.",
        "BUG-03 CONFIRMED — a live PENDING txn can be manually driven to FAILED with "
        "no guard; this is the state that lets Retry double-charge", CONFIRM),
    "TC-05": (
        "On the same transaction, changed Status to SUCCESS and saved. Rejected with: "
        "'Status SUCCESS may only be set by the Pinelabs transaction service, not "
        "manually.' Row did not become SUCCESS.",
        "PASS — the SUCCESS guard works (only SUCCESS is protected; FAILED/CANCELLED "
        "are not — see TC-04)", PASS),
    "TC-06": (
        "Opened the submitted unpaid invoice, clicked 'Pay on Machine'; the dialog "
        "opened listing PineLabs - Card / UPI. Selected 'PineLabs - Card' and clicked "
        "Start Payment. Server returned success=true, status=PENDING, response_code=0 "
        "('TXN ACCEPTED FOR PROCESSING'); a PENDING txn was created. The account "
        "pre-flight did NOT fire because config changed mid-session: PineLabs - Card is "
        "now mapped to Cash - PLM and the simulator endpoint (127.0.0.1:5000) is ACTIVE. "
        "Test txn was cancelled + deleted afterward.",
        "OBSERVED — happy-path terminal initiation works end-to-end vs the simulator. "
        "Pre-flight guard not exercised (accounts now mapped), so this run did not test "
        "the fail-fast path.", INFO),
    "TC-07": (
        "Ran int(float(amount)*100) for representative amounts: 262.15->26214, "
        "8.20->819, 1.15->114. Sweep: 9,174 of 200,000 two-decimal amounts undercharge "
        "by a paisa (4.59%).",
        "BUG-06 CONFIRMED (backend) — truncation, not rounding; plural.py already uses "
        "round()", CONFIRM),
    "TC-08": (
        "Fed a real 200 response with an HTML body to PineLabsClient._parse_response: "
        "returned success=True, data={}. Traced terminal.initiate_terminal_payment: "
        "response_code=None falls into the 'Synchronous decline' branch -> mark_failed.",
        "BUG-01 CONFIRMED (backend) — a malformed 200 marks a live payment FAILED with "
        "no reference to recover it", CONFIRM),
    "TC-09": (
        "Created Customer 'Acme Traders' (mobile 9000000001) and Supplier 'Acme Traders' "
        "(9111111111). With the map pinned to Customer, called _resolve_value on a doc "
        "with party_type=Supplier -> returned 9000000001 (the Customer's number). Demo "
        "parties removed.",
        "BUG-14 CONFIRMED (backend) — Dynamic Link ignored; wrong party's contact "
        "returned", CONFIRM),
}

# Bugs sheet Status updates from this run.
BUG_STATUS = {
    "BUG-01": "CONFIRMED (backend TC-08)",
    "BUG-03": "CONFIRMED (UI TC-04 + backend)",
    "BUG-06": "CONFIRMED (backend TC-07)",
    "BUG-10": "REFUTED (live test — check works)",
    "BUG-13": "REFUTED (UI TC-03 — buttons render on draft)",
    "BUG-14": "CONFIRMED (backend TC-09)",
}

wb = openpyxl.load_workbook(FILE)

# Test Cases: columns H (8)=Actual, I (9)=Verdict.
ws = wb["Test Cases"]
tc_col = 1
for r in range(2, ws.max_row + 1):
    tc = ws.cell(row=r, column=tc_col).value
    if tc in ACTUAL:
        actual, verdict, fill = ACTUAL[tc]
        ws.cell(row=r, column=8, value=actual).alignment = WRAP
        vcell = ws.cell(row=r, column=9, value=verdict)
        vcell.alignment = WRAP
        vcell.fill = fill
        for c in (8, 9):
            ws.cell(row=r, column=c).border = THIN

# Bugs: update Status (col 9).
wsb = wb["Bugs"]
for r in range(2, wsb.max_row + 1):
    bid = wsb.cell(row=r, column=1).value
    if bid in BUG_STATUS:
        wsb.cell(row=r, column=9, value=BUG_STATUS[bid]).alignment = WRAP

# Run Info sheet.
if "Run Info" in wb.sheetnames:
    del wb["Run Info"]
wsi = wb.create_sheet("Run Info")
info = [
    ["Field", "Value"],
    ["Site", "pine.local  (served at http://127.0.0.1:8015, default_site)"],
    ["Tool", "Playwright 1.61.1 (Chromium, headless) driving the Frappe Desk UI as Administrator"],
    ["Auth", "Server-minted Desk session (sid cookie); no password guessed or changed"],
    ["App code changed?", "NO — observation only, per instruction"],
    ["UI test cases", "TC-01..TC-06 run via Playwright against the live Desk"],
    ["Backend test cases", "TC-07..TC-09 run via the app's own functions (no gateway needed)"],
    ["Env change noted", "Mid-session, PineLabs - Card/UPI became mapped to Cash - PLM and the "
                          "simulator endpoint 127.0.0.1:5000 became ACTIVE. This is why TC-06 ran "
                          "the happy path instead of the account pre-flight."],
    ["Headline results", "BUG-03 confirmed in the UI (PENDING->FAILED accepted). BUG-01/06/14 "
                          "confirmed via backend. BUG-13 and BUG-10 REFUTED. SUCCESS guard (TC-05) works."],
    ["Test data left", "ACC-SINV-2026-00001 (Paid, real prior payment), -00002 (Draft), -00003 "
                        "(submitted unpaid). PLT-2026-00002 left as FAILED by TC-04. All test "
                        "txns on -00003 cleaned."],
]
for row in info:
    wsi.append(row)
for c in range(1, 3):
    cell = wsi.cell(row=1, column=c)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
wsi.column_dimensions["A"].width = 20
wsi.column_dimensions["B"].width = 110
for r in range(1, wsi.max_row + 1):
    wsi.cell(row=r, column=2).alignment = WRAP
    for c in (1, 2):
        wsi.cell(row=r, column=c).border = THIN

wb.save(FILE)
print("Recorded actuals + verdicts into", FILE)
print("Sheets:", wb.sheetnames)
