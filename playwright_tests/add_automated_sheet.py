# -*- coding: utf-8 -*-
"""(Re)build the 'Playwright Automated' sheet in pinelabs_test_cases.xlsx from the
executed UI suite, and stamp Actual Result + Status onto matching manual rows.
Update the A[] rows + timings after each run.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILE = "pinelabs_test_cases.xlsx"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
PASS_FONT = Font(bold=True, color="006100")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CTR = Alignment(horizontal="center", vertical="top")

# auto_id, manual_ref, module, title, type, prio, precondition, steps, data,
# expected, actual, status, time, spec_loc
A = [
 ("AUTO-AUTH","—","Auth","Session authenticates as Administrator","Setup","High",
  "Dev server up on 127.0.0.1:8015; storageState.json holds a server-minted sid cookie.",
  "1. GET /api/method/frappe.auth.get_logged_user with the session cookie.\n2. Read the body.",
  "storageState.json (minted sid).",
  "Body contains 'Administrator' (not redirected to /login).",
  "Returned {\"message\":\"Administrator\"}.","PASS","0.4s","pinelabs_ui.spec.mjs:79"),

 ("AUTO-BTN-01","TC-BTN-01","Buttons","Both buttons on a submitted unpaid invoice","Happy","High",
  "Sales Invoice in Payable Documents (both buttons enabled); a submitted invoice, outstanding>0.",
  "1. Navigate to /app/sales-invoice/<inv_clean>.\n2. Wait for form load + button injection.\n3. Query toolbar for 'Pay on Machine' and 'Send Payment Link'.",
  "inv_clean: a fresh submitted unpaid invoice (outstanding 1000).",
  "Both buttons present.","Both buttons rendered.","PASS","8.6s","pinelabs_ui.spec.mjs:87"),

 ("AUTO-BTN-02","TC-BTN-02","Buttons","No buttons on a fully paid invoice","Happy","High",
  "A fully paid Sales Invoice (outstanding 0).",
  "1. Navigate to /app/sales-invoice/<inv_paid>.\n2. Wait for load.\n3. Query for the two buttons.",
  "inv_paid = ACC-SINV-2026-00001 (Paid).",
  "Neither button present.","Both buttons absent.","PASS","8.0s","pinelabs_ui.spec.mjs:93"),

 ("AUTO-BTN-03","TC-BTN-03","Buttons","Buttons render on a draft (show_on_docstatus=0,1)","Edge","Medium",
  "Payable row show_on_docstatus='0,1'; a saved draft invoice.",
  "1. Navigate to /app/sales-invoice/<inv_draft>.\n2. Wait for load.\n3. Query for buttons.",
  "inv_draft = ACC-SINV-2026-00002 (Draft, outstanding populated).",
  "Both buttons render (BUG-13 refuted).","Both buttons rendered on the draft.","PASS","7.1s","pinelabs_ui.spec.mjs:99"),

 ("AUTO-BTN-05","TC-BTN-05","Buttons","No buttons on a cancelled invoice","Negative","Medium",
  "A cancelled Sales Invoice (docstatus 2).",
  "1. Navigate to /app/sales-invoice/<inv_cancelled>.\n2. Wait for load.\n3. Query for buttons.",
  "inv_cancelled: a submitted-then-cancelled invoice.",
  "Neither button present.","Both buttons absent.","PASS","6.7s","pinelabs_ui.spec.mjs:106"),

 ("AUTO-BTN-07","TC-BTN-07","Buttons","No buttons on a doctype without a Payable row","Edge","Low",
  "A doctype (Customer) with no Payable Documents row.",
  "1. Navigate to /app/customer/<record>.\n2. Wait for load.\n3. Query for the PineLabs buttons.",
  "Customer 'Pine Demo Customer'.",
  "No PineLabs buttons.","Both buttons absent.","PASS","7.0s","pinelabs_ui.spec.mjs:112"),

 ("AUTO-SM-05","TC-SM-05","State machine","Changing flow_type after creation is rejected","Negative","Medium",
  "An existing transaction (flow_type Terminal).",
  "1. Open the transaction form.\n2. Change flow_type Select to Plural.\n3. Ctrl+S.\n4. Read error.",
  "txn_reject.",
  "Rejected: 'flow_type is locked after creation'.","Rejected with the lock message.","PASS","8.8s","pinelabs_ui.spec.mjs:126"),

 ("AUTO-SM-04","TC-SM-04","State machine","Status field read-only; manual FAILED no effect (BUG-03 FIXED)","Negative","High",
  "A PENDING transaction.",
  "1. Open /app/pinelabs-transaction/<txn_fail>.\n2. Assert the Status field is read-only (cur_frm meta).\n3. Force a Select change + Ctrl+S.\n4. Re-read status via REST.",
  "txn_fail (PENDING).",
  "Status field is read-only AND the forced edit does not take effect; status stays PENDING.",
  "read-only=true; status stayed PENDING (BUG-03 fixed).","PASS (BUG-03 fixed)","6.8s","pinelabs_ui.spec.mjs:131"),

 ("AUTO-SM-06","TC-SM-06","State machine","Deleting a transaction linked to a Payment Entry is blocked","Negative","Medium",
  "A SUCCESS transaction linked to a submitted Payment Entry.",
  "1. Open the transaction form.\n2. Trigger delete.\n3. Assert it throws and the record still exists.",
  "txn_with_pe = PLT-2026-00001 (SUCCESS, linked to ACC-PAY-2026-00001).",
  "Delete blocked (linked to a Payment Entry); record present.","Delete threw; record still exists.","PASS","6.2s","pinelabs_ui.spec.mjs:145"),

 ("AUTO-TERM-03","TC-TERM-03","Terminal","Start Payment with no mode does not proceed","Negative","Medium",
  "A submitted unpaid invoice (Pay on Machine enabled).",
  "1. Open the invoice.\n2. Click 'Pay on Machine'.\n3. Without a mode, invoke the dialog primary action.\n4. Observe get_values() null and dialog still open.",
  "inv_clean.",
  "Does not proceed; dialog stays open; no transaction.","primary action 'no-values'; dialog open.","PASS","12.6s","pinelabs_ui.spec.mjs:168"),

 ("AUTO-TERM-11","TC-TERM-11","Terminal","Second attempt while one in progress is refused","Negative","High",
  "An invoice with an active (PENDING) transaction.",
  "1. Open the invoice.\n2. Pay on Machine → PineLabs - Card → Start.\n3. Capture the initiate response.",
  "inv_active with an active txn.",
  "Refused: response indicates a payment already in progress.","Response contained 'already in progress'.","PASS","34.9s","pinelabs_ui.spec.mjs:175"),

 ("AUTO-TERM-01","TC-TERM-01","Terminal","Pay on Machine (Card) creates a transaction","Happy","High",
  "A submitted unpaid invoice; mode account mapped.",
  "1. Open the invoice.\n2. Pay on Machine → pick Card → Start.\n3. Query REST for a Terminal transaction.",
  "inv_clean; Local Simulator Machine.",
  "The UI->server->transaction flow runs; a Terminal transaction is created (gateway accept/decline depends on the terminal/simulator).",
  "Transaction created.","PASS","34.5s","pinelabs_ui.spec.mjs:185"),

 ("AUTO-PLURAL-06","TC-PLURAL-06","Payment Link","Send Payment Link on a draft is rejected","Negative","Medium",
  "A draft invoice with the Send Payment Link button visible (show_on_docstatus=0,1).",
  "1. Open /app/sales-invoice/<inv_draft>.\n2. Click 'Send Payment Link'.\n3. Capture the initiate_payment_link response.",
  "inv_draft (Draft).",
  "Rejected: invoice must be submitted to send a payment link.","Response contained 'must be submitted'.","PASS","11.0s","pinelabs_ui.spec.mjs:197"),

 ("AUTO-CG-01","TC-CG-01","Conflict guard","Manual Payment Entry blocked while a payment is active","Negative","High",
  "An invoice with an active (INITIATED/PENDING) transaction.",
  "1. Open the invoice.\n2. Build a Payment Entry (get_payment_entry) and attempt to save.\n3. Read server messages.",
  "inv_active with an active txn.",
  "Save refused by the conflict guard (a Pine Labs payment is in progress).","Save blocked; message contained 'in progress'.","PASS","7.2s","pinelabs_ui.spec.mjs:213"),
]

wb = openpyxl.load_workbook(FILE)
if "Playwright Automated" in wb.sheetnames:
    del wb["Playwright Automated"]
ws = wb.create_sheet("Playwright Automated")
cols = ["Auto TC ID","Manual Ref","Module","Title","Type","Priority","Preconditions",
        "Automation Steps","Test Data / Fixtures","Expected Result","Actual Result",
        "Status","Exec Time","Spec Location"]
ws.append([""]*len(cols))  # banner row 1
ws.append(cols)            # header row 2
for c in range(1, len(cols)+1):
    cell = ws.cell(2, c); cell.fill = HEAD_FILL; cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = THIN
ws.row_dimensions[2].height = 28
ws.freeze_panes = "A3"
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
b = ws.cell(1,1); b.value = ("Executed Playwright UI suite — 14/14 PASSED (post Priority-1 fixes; TC-SM-04 now "
    "verifies read-only status). Drives the Frappe Desk as a "
    "user via a server-minted session. Run: playwright_tests/ (setup_fixtures.py -> npx playwright test).")
b.font = Font(name="Calibri", bold=True, color="1F4E78"); b.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 30

for row in A:
    ws.append(list(row))
    r = ws.max_row
    for c in range(1, len(cols)+1):
        cell = ws.cell(r, c); cell.border = THIN
        cell.alignment = CTR if c in (5,6,12,13) else WRAP
    sc = ws.cell(r, 12); sc.fill = PASS_FILL; sc.font = PASS_FONT

widths = [15,12,15,32,10,9,34,52,28,46,40,22,10,24]
for i,w in enumerate(widths,1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Stamp manual sheet (cols 10 Actual, 11 Status).
manual_map = {r[1]: (r[10], r[11]) for r in A if r[1] != "—"}
tcs = wb["Test Cases"]
for rr in range(2, tcs.max_row+1):
    tcid = tcs.cell(rr, 1).value
    if tcid in manual_map:
        actual, status = manual_map[tcid]
        ac = tcs.cell(rr, 10, value=f"[Automated] {actual}"); ac.alignment = WRAP
        st = tcs.cell(rr, 11, value="Pass"); st.alignment = CTR; st.fill = PASS_FILL; st.font = PASS_FONT

wb.save(FILE)
print(f"Rebuilt 'Playwright Automated' with {len(A)} rows; stamped {len(manual_map)} manual rows.")
