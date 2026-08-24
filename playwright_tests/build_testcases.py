# -*- coding: utf-8 -*-
"""Build pinelabs_test_cases.xlsx — full manual/functional test-case suite.

Sheets: Test Cases (all TCs), Summary (counts), Legend (how to use).
Expected Result = correct/intended behavior. Where the app's ACTUAL behavior
currently deviates, the Expected cell ends with 'KNOWN ISSUE: BUG-xx'.
Actual Result + Status columns are left blank for the tester to fill on execution.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "pinelabs_test_cases.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
MOD_FILL = PatternFill("solid", fgColor="DDEBF7")
MOD_FONT = Font(bold=True, size=11, color="1F4E78")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")

TYPE_FILL = {
    "Happy": "C6EFCE", "Negative": "FFC7CE", "Edge": "FFEB9C",
    "Boundary": "BDD7EE", "Security": "E4DFEC", "Concurrency": "D9F2F0",
}
PRIO_FILL = {"High": "FF7C80", "Medium": "FFD966", "Low": "A9D08E"}

# Each row: (id, module, title, type, priority, precondition, steps, data, expected)
T = []
def add(_id, mod, title, typ, prio, pre, steps, data, exp):
    T.append((_id, mod, title, typ, prio, pre, steps, data, exp))

# ── A. INSTALLATION & SEED ───────────────────────────────────────────────
add("SETUP-01","Install & Seed","Fresh install seeds 3 Modes of Payment","Happy","High",
    "Clean site without pinelabs_fin.",
    "1. Install app: bench --site <s> install-app pinelabs_fin.\n2. Open Mode of Payment list.",
    "—",
    "'PineLabs - Card', 'PineLabs - UPI', 'PineLabs - Payment Link' exist, all Enabled, type Bank.")
add("SETUP-02","Install & Seed","Upgrade migrate seeds modes idempotently","Happy","Medium",
    "Existing site upgraded from a version without the seed.",
    "1. bench --site <s> migrate.\n2. Check Mode of Payment list.",
    "—",
    "The 3 canonical modes are created if missing; no duplicates on re-run.")
add("SETUP-03","Install & Seed","Re-running migrate does not duplicate modes","Edge","Medium",
    "Modes already seeded.",
    "1. Run bench migrate twice.\n2. Count PineLabs modes.",
    "—",
    "Exactly 3 PineLabs modes exist after repeated migrations.")
add("SETUP-04","Install & Seed","Renaming a canonical mode breaks integration","Negative","Medium",
    "Modes seeded; a payable doctype configured.",
    "1. Rename 'PineLabs - Card' to something else.\n2. Attempt Pay on Machine with Card.",
    "Rename via Mode of Payment form.",
    "Integration fails to resolve the mode (documented limitation — names are the contract). Restore the name to fix.")

# ── B. SETTINGS VALIDATION ───────────────────────────────────────────────
add("SET-01","Settings","Enable integration + one machine auto-defaults","Happy","High",
    "Pinelabs Settings open.",
    "1. Tick Enabled.\n2. Add a single Machine row with all credentials.\n3. Save.",
    "Client ID, Security Token, Merchant ID, Store ID, Terminal ID.",
    "Saved; the single machine is auto-marked Is Default.")
add("SET-02","Settings","Two machines both Is Default rejected","Negative","High",
    "Two machine rows exist.",
    "1. Tick Is Default on both.\n2. Save.",
    "2 machine rows.",
    "Error: exactly one machine may be Is Default.")
add("SET-03","Settings","Enable Payment Links requires Plural credentials","Negative","High",
    "Payment Links section visible.",
    "1. Tick Enable Payment Links.\n2. Leave Plural Client ID blank.\n3. Save.",
    "—",
    "Mandatory error on Plural Client ID / Secret / Webhook Secret.")
add("SET-04","Settings","Webhook Secret below 16 chars rejected","Boundary","High",
    "Enable Payment Links ticked, Plural creds filled.",
    "1. Enter a 3-char Webhook Secret.\n2. Save.",
    "Webhook Secret = 'abc'.",
    "Error: Webhook Secret must be at least 16 characters.")
add("SET-05","Settings","Webhook Secret exactly 16 chars accepted","Boundary","Medium",
    "Enable Payment Links ticked, Plural creds filled.",
    "1. Enter a 16-char Webhook Secret.\n2. Save.",
    "Webhook Secret = 16 chars.",
    "Saved successfully.")
add("SET-06","Settings","Blank Plural Base URL auto-fills sandbox","Happy","Medium",
    "Enable Payment Links ticked.",
    "1. Leave Plural API Base URL blank.\n2. Save.\n3. Re-open.",
    "—",
    "Base URL auto-filled to the sandbox host (pluraluat...).")
add("SET-07","Settings","Base URL trailing slash trimmed","Edge","Low",
    "Enable Payment Links ticked.",
    "1. Enter base URL with a trailing '/'.\n2. Save.",
    "https://host/",
    "Stored without the trailing slash.")
add("SET-08","Settings","Test Plural Connection success","Happy","Medium",
    "Valid Plural credentials.",
    "1. Click Test Plural Connection.",
    "Valid client id/secret.",
    "Success dialog: token generated (expires ~60 min).")
add("SET-09","Settings","Test Plural Connection with bad creds","Negative","Medium",
    "Invalid Plural credentials.",
    "1. Enter wrong client secret.\n2. Click Test Plural Connection.",
    "Wrong secret.",
    "Failure message; token not obtained.")
add("SET-10","Settings","Changing credentials busts cached token","Edge","Low",
    "Token previously cached.",
    "1. Change Plural Client ID.\n2. Save.\n3. Click Test Connection.",
    "New client id.",
    "Fresh token fetched (old cache cleared).")
add("SET-11","Settings","Payable Doctype missing Amount Field rejected","Negative","High",
    "Payable Documents section.",
    "1. Add a row with a DocType but no Amount Field.\n2. Save.",
    "DocType = Sales Invoice.",
    "Error: Amount Field required.")
add("SET-12","Settings","Duplicate Payable Doctype rows rejected","Negative","Medium",
    "Payable Documents section.",
    "1. Add two rows for the same DocType.\n2. Save.",
    "Two 'Sales Invoice' rows.",
    "Error: duplicate Payable Doctype row.")
add("SET-13","Settings","Create Payment Entry off requires status mapping","Negative","Medium",
    "Payable Documents row.",
    "1. Untick Create Payment Entry.\n2. Leave Status Field / Paid Status Value blank.\n3. Save.",
    "—",
    "Error: Status Field and Paid Status Value required.")
add("SET-14","Settings","Contact mapping Direct Field without field name","Negative","Low",
    "Payable Documents row.",
    "1. Set Mobile = Direct Field.\n2. Leave Field Name blank.\n3. Save.",
    "—",
    "Error: Field Name empty for Direct Field mapping.")
add("SET-15","Settings","Contact mapping Linked DocType missing pieces","Negative","Low",
    "Payable Documents row.",
    "1. Set Email = Linked DocType.\n2. Leave Link Field / Linked DocType / Field blank.\n3. Save.",
    "—",
    "Error listing the missing Linked DocType pieces.")

# ── C. BUTTON RENDERING ──────────────────────────────────────────────────
add("BTN-01","Buttons","Both buttons on submitted unpaid invoice","Happy","High",
    "Sales Invoice in Payable Documents (both buttons enabled), submitted, outstanding > 0.",
    "1. Open the submitted invoice.",
    "outstanding_amount = 1000.",
    "'Pay on Machine' and 'Send Payment Link' buttons visible.")
add("BTN-02","Buttons","No buttons on fully-paid invoice","Happy","High",
    "Invoice fully paid (outstanding 0).",
    "1. Open the paid invoice.",
    "outstanding_amount = 0.",
    "Neither button appears.")
add("BTN-03","Buttons","Buttons on a draft invoice (show_on_docstatus incl 0)","Edge","Medium",
    "Payable row show_on_docstatus='0,1'; a saved draft invoice.",
    "1. Open the draft invoice.",
    "Draft, outstanding_amount populated (~grand_total).",
    "Both buttons render on the draft (outstanding is populated on save).")
add("BTN-04","Buttons","Draft hidden when show_on_docstatus='1'","Edge","Medium",
    "Payable row show_on_docstatus='1'; a draft invoice.",
    "1. Open the draft.",
    "—",
    "Buttons hidden (docstatus gate).")
add("BTN-05","Buttons","Cancelled invoice shows no buttons","Negative","Medium",
    "A cancelled invoice.",
    "1. Open the cancelled invoice.",
    "docstatus = 2.",
    "No buttons.")
add("BTN-06","Buttons","Only enabled button shows","Happy","Low",
    "Payable row with only Pay on Machine enabled.",
    "1. Open a qualifying invoice.",
    "—",
    "Only 'Pay on Machine' appears; 'Send Payment Link' hidden.")
add("BTN-07","Buttons","Doctype not configured → no buttons","Edge","Low",
    "A doctype with no Payable Documents row.",
    "1. Open a record of that doctype.",
    "—",
    "No PineLabs buttons.")

# ── D. TERMINAL (Pay on Machine) ─────────────────────────────────────────
add("TERM-01","Terminal","Card payment success end-to-end","Happy","High",
    "Machine configured & enabled; PineLabs - Card mapped to an account for the company; submitted unpaid invoice.",
    "1. Click Pay on Machine.\n2. Select PineLabs - Card, Start Payment.\n3. Customer taps card on terminal.\n4. Wait for dialog to settle.",
    "Card tap on simulator/terminal.",
    "Dialog shows success; invoice flips to Paid; a submitted Payment Entry is created (mode PineLabs - Card, reference_no = Plutus id).")
add("TERM-02","Terminal","UPI payment success","Happy","High",
    "As TERM-01 with UPI mapped.",
    "1. Pay on Machine → PineLabs - UPI → Start.\n2. Customer scans UPI QR.",
    "UPI scan.",
    "Invoice Paid; Payment Entry created with mode PineLabs - UPI.")
add("TERM-03","Terminal","Mode of Payment required","Negative","Medium",
    "Pay on Machine dialog open.",
    "1. Click Start Payment without selecting a mode.",
    "—",
    "Validation stops; 'Mode of Payment is required'. No transaction created.")
add("TERM-04","Terminal","No machine configured","Negative","High",
    "Integration enabled but no enabled/default machine.",
    "1. Click Pay on Machine → select mode → Start.",
    "—",
    "Error: no machine resolves; configure a machine / default. No gateway call.")
add("TERM-05","Terminal","Mode account not mapped → fail-fast","Negative","High",
    "Chosen mode has NO default account for the invoice company; PE will be created.",
    "1. Pay on Machine → select mode → Start.",
    "Unmapped PineLabs - Card.",
    "Clean error 'Mode of Payment has no default account…' BEFORE any gateway call; customer never charged.")
add("TERM-06","Terminal","Terminal declines synchronously (RC 2)","Negative","High",
    "Terminal set to decline.",
    "1. Pay on Machine → Start.\n2. Gateway returns ResponseCode 2.",
    "Simulated decline.",
    "Transaction FAILED; dialog shows failure; Retry button offered.")
add("TERM-07","Terminal","Terminal timeout (RC 1006)","Edge","Medium",
    "Poll returns 1006.",
    "1. Start a payment.\n2. Terminal does not respond; poll returns 1006.",
    "Simulated timeout.",
    "Transaction FAILED immediately (definitive timeout).")
add("TERM-08","Terminal","Voided at terminal (RC 1008)","Edge","Medium",
    "Poll returns 1008.",
    "1. Start payment.\n2. Void at the terminal UI.",
    "Simulated void.",
    "Transaction CANCELLED; dialog closes.")
add("TERM-09","Terminal","Still-pending past cutoff → age-out","Edge","Medium",
    "Poll keeps returning 1001 beyond auto-cancel minutes.",
    "1. Start payment.\n2. Leave it beyond cutoff (default 5 min).",
    "auto_cancel_duration_in_minutes.",
    "App force-cancels, does a final poll; SUCCESS if customer paid in the race window, else FAILED (Aged out).")
add("TERM-10","Terminal","Close dialog → cron reconciles","Edge","High",
    "A PENDING terminal payment; scheduler enabled.",
    "1. Start payment.\n2. Close the browser tab before it settles.\n3. Wait ~1 minute.",
    "—",
    "The 1-minute cron polls the gateway and finalizes the transaction (SUCCESS/FAILED) without the browser.")
add("TERM-11","Terminal","Second attempt while one in flight","Negative","High",
    "A PENDING/INITIATED terminal txn exists for the invoice.",
    "1. On the same invoice, click Pay on Machine again → Start.",
    "—",
    "Error: 'A Terminal payment (PLT-…) is already in progress'. No second charge.")
add("TERM-12","Terminal","Cancel while customer paid in race","Edge","High",
    "Payment in flight; customer taps at the same instant as Cancel.",
    "1. Start payment.\n2. Click Cancel just as the customer taps.\n3. Gateway rejects the void; re-poll shows APPROVED.",
    "Race window.",
    "Transaction marked SUCCESS (not CANCELLED); PE created. KNOWN ISSUE: BUG-02-cancel-UI — the modal ignores the SUCCESS response and reports cancel (see BUG list; distinct from the Plural BUG-02).")
add("TERM-13","Terminal","User cancels an in-flight payment","Happy","Medium",
    "A PENDING terminal txn, customer has not paid.",
    "1. Click Cancel in the waiting dialog.",
    "—",
    "Gateway cancel sent; transaction CANCELLED.")
add("TERM-14","Terminal","Retry after failure starts a new payment","Edge","Medium",
    "A FAILED terminal txn with the dialog showing Retry.",
    "1. Click Retry.",
    "—",
    "A brand-new terminal payment initiates for the same invoice.")
add("TERM-15","Terminal","Draft invoice auto-submits on success","Boundary","Medium",
    "Payable row allows drafts; a draft Sales Invoice.",
    "1. Pay on Machine on the draft → success.",
    "Draft SI.",
    "On success the invoice is submitted, then the Payment Entry is created and it becomes Paid.")
add("TERM-16","Terminal","Cancelled source doc rejected","Negative","Medium",
    "A cancelled invoice.",
    "1. Attempt Pay on Machine (via API if button hidden).",
    "docstatus 2.",
    "Error: cancelled doc cannot take a payment.")
add("TERM-17","Terminal","Accepted but no Plutus reference → fail-fast","Edge","Low",
    "Gateway returns ResponseCode 0 but no PlutusTransactionReferenceID.",
    "1. Start payment.",
    "Simulated missing ref.",
    "Transaction FAILED with a clear message (cannot track).")
add("TERM-18","Terminal","create_payment_entry=0 override skips PE","Edge","Medium",
    "API called with create_payment_entry=0 on a PE-native doctype.",
    "1. Call initiate_terminal_payment(..., create_payment_entry=0) and complete.",
    "Override 0.",
    "Transaction SUCCESS; NO Payment Entry created.")
add("TERM-19","Terminal","Malformed non-JSON 200 from gateway","Edge","High",
    "Gateway/proxy returns HTTP 200 with a non-JSON body.",
    "1. Start payment; gateway returns 200 + HTML.",
    "Simulated proxy page.",
    "Treated as a failed call with a clear 'unreadable response' error (not a phantom decline). FIXED (BUG-01).")

# ── E. PLURAL (Send Payment Link) ────────────────────────────────────────
add("PLURAL-01","Payment Link","Generate all-methods link","Happy","High",
    "Payment Links enabled; Plural creds valid; PineLabs - Payment Link account mapped; submitted invoice with email/mobile.",
    "1. Click Send Payment Link.\n2. Confirm the green success popup with the link.",
    "Customer email + mobile present.",
    "Link generated; Pine Labs sends SMS + email; a PENDING Plural transaction is created.")
add("PLURAL-02","Payment Link","Restricted flow (card/upi/nb only)","Happy","Medium",
    "As PLURAL-01.",
    "1. Call initiate_payment_link(flow='upi_only').",
    "flow=upi_only.",
    "Hosted page restricted to UPI; payment_method preset accordingly.")
add("PLURAL-03","Payment Link","Customer pays → webhook finalizes","Happy","High",
    "A PENDING Plural link; webhook URL registered.",
    "1. Customer completes payment on the hosted page.\n2. Plural POSTs the webhook.",
    "Real/simulated payment.",
    "Invoice flips to Paid; Payment Entry created (mode PineLabs - Payment Link).")
add("PLURAL-04","Payment Link","No email/mobile → error with trace","Negative","High",
    "Invoice with no resolvable email or mobile; no inline override.",
    "1. Click Send Payment Link.",
    "No contact fields.",
    "Error: customer email/mobile required, with a resolver trace; no link sent.")
add("PLURAL-05","Payment Link","Inline email/mobile override","Happy","Medium",
    "Invoice without contact fields.",
    "1. Call initiate_payment_link(customer_email=…, customer_mobile=…).",
    "Inline contacts.",
    "Link generated using the inline values (resolver bypassed for those).")
add("PLURAL-06","Payment Link","Draft invoice rejected","Negative","Medium",
    "A draft invoice.",
    "1. Send Payment Link on the draft.",
    "docstatus 0.",
    "Error: invoice must be submitted to send a payment link.")
add("PLURAL-07","Payment Link","Custom expiry honored","Edge","Low",
    "Payment Links enabled.",
    "1. Call initiate_payment_link(expiry_minutes=60).",
    "expiry 60.",
    "Link expiry_date is ~60 minutes out; default is 1440 when omitted.")
add("PLURAL-08","Payment Link","Payment Link account not mapped → fail-fast","Negative","High",
    "PineLabs - Payment Link has no account for the company; PE will be created.",
    "1. Send Payment Link.",
    "Unmapped mode.",
    "Clean pre-flight error BEFORE contacting Plural.")
add("PLURAL-09","Payment Link","Missed webhook → cron reconciles","Edge","High",
    "A PENDING Plural link; webhook not delivered.",
    "1. Customer pays; webhook dropped.\n2. Wait ~1 minute.",
    "—",
    "The 1-minute cron pulls status and finalizes the transaction → invoice Paid.")
add("PLURAL-10","Payment Link","Second link after first failed uses unique reference","Edge","High",
    "First link FAILED/expired for an invoice.",
    "1. Send Payment Link again on the same invoice.\n2. Open both Pinelabs Transactions → Request Payload.",
    "—",
    "Each transaction's merchant_payment_link_reference equals its own PLT-… name (distinct), not the invoice name. (BUG-02 fixed.)")
add("PLURAL-11","Payment Link","Source doc deleted before finalize","Edge","Low",
    "A PENDING Plural link whose invoice was deleted.",
    "1. Delete the invoice.\n2. Run Refresh Status on the transaction.",
    "—",
    "Transaction marked FAILED with 'source no longer exists'; cron stops looping.")
add("PLURAL-12","Payment Link","Link creation fails at gateway","Negative","Medium",
    "Plural returns an error.",
    "1. Send Payment Link; gateway rejects.",
    "Bad creds / gateway error.",
    "Transaction FAILED; error shown; no link delivered.")

# ── F. WEBHOOK ───────────────────────────────────────────────────────────
add("WH-01","Webhook","Valid signed paid event → success","Happy","High",
    "Webhook secret configured; a PENDING Plural txn with matching order_id.",
    "1. POST ORDER_PROCESSED with a valid X-Verify HMAC over the raw body.",
    "Valid signature.",
    "HTTP 200; transaction SUCCESS; Payment Entry created.")
add("WH-02","Webhook","Missing X-Verify header → 401","Negative","High",
    "Webhook secret configured.",
    "1. POST without X-Verify.",
    "—",
    "HTTP 401; not processed.")
add("WH-03","Webhook","Invalid signature → 401","Negative","High",
    "Webhook secret configured.",
    "1. POST with a wrong X-Verify.",
    "Tampered body / wrong secret.",
    "HTTP 401; not processed.")
add("WH-04","Webhook","Secret unconfigured → fail-closed 401","Security","High",
    "No webhook secret set.",
    "1. POST any signed webhook.",
    "—",
    "HTTP 401 (fail-closed); never silently accepted.")
add("WH-05","Webhook","Duplicate event_id ignored (idempotent)","Edge","Medium",
    "An event already processed.",
    "1. POST the same event_id twice.",
    "Same event_id.",
    "Second delivery → HTTP 200 'Duplicate event ignored'; no double processing.")
add("WH-06","Webhook","Unknown order_id → 404","Negative","Medium",
    "No transaction matches the order.",
    "1. POST a valid signed webhook for an unknown order.",
    "order_id='no-such'.",
    "HTTP 404 (Plural won't retry an unrecoverable case).")
add("WH-07","Webhook","Already SUCCESS → idempotent 200","Edge","Medium",
    "Transaction already SUCCESS.",
    "1. Re-deliver the paid webhook.",
    "—",
    "HTTP 200 'Already processed'; no change.")
add("WH-08","Webhook","Non-POST method → 405","Negative","Low",
    "Endpoint reachable.",
    "1. GET the webhook URL.",
    "—",
    "HTTP 405 Only POST allowed.")
add("WH-09","Webhook","Failed processing then retry","Edge","High",
    "A signed paid webhook whose processing throws (e.g. txn already FAILED).",
    "1. POST it (processing errors).\n2. Plural re-delivers the same event_id.",
    "—",
    "The retry re-processes (dedup only skips terminally-processed events). FIXED (BUG-15).")

# ── G. PAYMENT ENTRY / FINALIZATION ──────────────────────────────────────
add("PE-01","Payment Entry","PE built for Sales Invoice on success","Happy","High",
    "Mapped account; a submitted SI paid via terminal/link.",
    "1. Complete a payment.\n2. Open the created Payment Entry.",
    "—",
    "PE via get_payment_entry: correct mode_of_payment, paid_amount = txn amount, reference_no = Plutus/Plural id; submitted.")
add("PE-02","Payment Entry","Idempotent — existing PE with same reference reused","Edge","High",
    "A PE already exists with reference_no = payment id.",
    "1. Trigger finalize again (webhook replay / cron).",
    "—",
    "No duplicate PE; the existing one is reused.")
add("PE-03","Payment Entry","Transaction already linked → no duplicate","Edge","Medium",
    "Transaction.payment_entry already set.",
    "1. Re-run mark_success.",
    "—",
    "Returns the linked PE; no new PE.")
add("PE-04","Payment Entry","Custom doctype create_payment_entry=1","Edge","Medium",
    "A non-PE-native doctype with a customer link; account mapped.",
    "1. Call the API with create_payment_entry=1 and complete.",
    "Custom doctype.",
    "A standalone Payment Entry is built manually (party=Customer, paid_to=mode account).")
add("PE-05","Payment Entry","Custom doctype no customer → pre-flight error","Negative","Medium",
    "A custom doctype with no resolvable customer; create_payment_entry=1.",
    "1. Initiate the payment.",
    "No customer / default.",
    "Pre-flight error BEFORE charging: cannot resolve a Customer.")
add("PE-06","Payment Entry","Mapped-field finalization (no PE)","Edge","Medium",
    "Payable row finalize_via_pe=0 with status_field + paid_status_value.",
    "1. Complete a payment on that doctype.",
    "—",
    "No PE; the configured status field is set to the paid value (and doc submitted if value='Submitted').")
add("PE-07","Payment Entry","Permissive path leaves source untouched","Edge","Low",
    "A doctype with no Payable row and no override.",
    "1. Drive a payment to SUCCESS.",
    "—",
    "Transaction SUCCESS; source doc not auto-updated (caller reacts to realtime event).")
add("PE-08","Payment Entry","Conflict — invoice already fully paid","Edge","High",
    "A manual PE was recorded while the Pine Labs payment was in flight.",
    "1. Let the Pine Labs payment finalize.",
    "outstanding = 0 externally.",
    "Transaction SUCCESS, payment_entry empty, error_message notes manual reconciliation needed (customer was charged).")

# ── H. STATE MACHINE ─────────────────────────────────────────────────────
add("SM-01","State Machine","Valid path INITIATED→PENDING→SUCCESS","Happy","High",
    "A transaction driven by the service.",
    "1. Initiate → PENDING → finalize SUCCESS.",
    "—",
    "All transitions allowed; ends SUCCESS with a PE.")
add("SM-02","State Machine","Manual UI edit to SUCCESS blocked","Negative","High",
    "A PENDING transaction open in the Desk.",
    "1. Change Status to SUCCESS.\n2. Save.",
    "—",
    "Rejected: 'Status SUCCESS may only be set by the Pinelabs transaction service'.")
add("SM-03","State Machine","Illegal transition PENDING→INITIATED blocked","Negative","Medium",
    "A PENDING transaction.",
    "1. Set Status back to INITIATED.\n2. Save.",
    "—",
    "Rejected as an illegal transition.")
add("SM-04","State Machine","Manual edit PENDING→FAILED","Negative","High",
    "A PENDING transaction open in the Desk.",
    "1. Change Status to FAILED.\n2. Save.",
    "—",
    "Rejected (only the service may transition); the status field is also read-only in the UI. FIXED (BUG-03).")
add("SM-05","State Machine","flow_type locked after creation","Negative","Medium",
    "An existing transaction.",
    "1. Change flow_type Terminal↔Plural.\n2. Save.",
    "—",
    "Rejected: flow_type is locked after creation.")
add("SM-06","State Machine","Delete transaction linked to a PE blocked","Negative","Medium",
    "A SUCCESS transaction with a linked PE.",
    "1. Try to delete it.",
    "—",
    "Rejected: linked to a Payment Entry.")
add("SM-07","State Machine","Duplicate provider id on active row blocked","Edge","Medium",
    "Two active transactions.",
    "1. Set the same payment_id/order_id on a second active row.\n2. Save.",
    "—",
    "Rejected: two active transactions cannot share a provider id.")
add("SM-08","State Machine","Concurrent finalize is serialized","Concurrency","High",
    "Webhook + cron + foreground poll race to finalize one transaction.",
    "1. Trigger two finalizers on the same PENDING row simultaneously.",
    "—",
    "Row lock serializes them; only one PE; the other no-ops (idempotent).")

# ── I. DOUBLE-SUBMIT / CONCURRENCY ───────────────────────────────────────
add("CONC-01","Concurrency","Rapid double-click on Start Payment","Concurrency","High",
    "Submitted invoice, no active txn.",
    "1. Double-click Start Payment quickly (or fire two initiate calls together).",
    "—",
    "Only one transaction/charge; the second concurrent attempt is refused 'already in progress'. FIXED (BUG-08).")

# ── J. ROUTING ───────────────────────────────────────────────────────────
add("ROUTE-01","Routing","Single Machine mode uses default","Happy","Medium",
    "Routing mode = Single Machine; one default machine.",
    "1. Initiate a terminal payment.",
    "—",
    "The default machine is used.")
add("ROUTE-02","Routing","Mapping by user","Happy","Medium",
    "Routing = Mapping; a rule maps the logged-in user to a machine.",
    "1. As that user, initiate a payment.",
    "User mapping.",
    "The mapped machine is chosen.")
add("ROUTE-03","Routing","Mapping by warehouse/POS profile","Happy","Low",
    "Routing = Mapping; a warehouse/POS-profile rule.",
    "1. Initiate on a doc carrying that warehouse/POS profile.",
    "—",
    "The mapped machine is chosen from the doc context.")
add("ROUTE-04","Routing","No mapping match → default fallback","Edge","Medium",
    "Routing = Mapping; no rule matches the context.",
    "1. Initiate.",
    "—",
    "Falls back to the default machine.")
add("ROUTE-05","Routing","Multiple matches → alphabetical first","Edge","Low",
    "Two rules match the context.",
    "1. Initiate.",
    "Two matching machines.",
    "The alphabetically-first machine name is chosen (deterministic).")

# ── K. CUSTOMER CONTACT RESOLUTION ───────────────────────────────────────
add("CC-01","Contact","Direct field mapping","Happy","Medium",
    "Payable row maps mobile/email to direct fields.",
    "1. Send Payment Link.",
    "contact_email on the doc.",
    "Configured direct fields are used for the customer.")
add("CC-02","Contact","Linked DocType mapping","Happy","Medium",
    "Payable row maps mobile via a linked doctype (e.g. Address.phone).",
    "1. Send Payment Link.",
    "customer_address → phone.",
    "Value read from the linked doctype.")
add("CC-03","Contact","Auto-detect common fields","Edge","Low",
    "No mapping configured.",
    "1. Send Payment Link on a doc with mobile_no/email_id.",
    "—",
    "App auto-detects common field names.")
add("CC-04","Contact","Mobile normalization","Boundary","Low",
    "Mobile has country code / extra digits.",
    "1. Send Payment Link with mobile '+91 98765 43210'.",
    "12-digit with 91 prefix.",
    "Normalized to the 10-digit number with country_code 91.")
add("CC-05","Contact","Dynamic Link party resolves correct entity","Edge","Medium",
    "A doc with party_type/party (Dynamic Link) mapped to a linked doctype.",
    "1. Send Payment Link where party_type=Supplier but map pins Customer.",
    "Same name as Customer + Supplier.",
    "Correct: resolve the actual party. KNOWN ISSUE: BUG-14 — returns the pinned doctype's record (wrong party's contact).")

# ── L. CONFLICT GUARD ────────────────────────────────────────────────────
add("CG-01","Conflict Guard","Manual PE blocked while Pine Labs active","Negative","High",
    "A Pinelabs Transaction INITIATED/PENDING for the invoice.",
    "1. Create a manual Payment Entry against that invoice.\n2. Save.",
    "—",
    "Rejected: a Pine Labs payment is in progress; finish/cancel it first.")
add("CG-02","Conflict Guard","App's own PE bypasses the guard","Happy","Medium",
    "A Pine Labs payment finalizing (its own txn still PENDING at insert).",
    "1. Let mark_success insert its PE.",
    "—",
    "Internal PE (flags.pinelabs_internal_pe) inserts without being blocked.")
add("CG-03","Conflict Guard","Manual PE allowed after txn terminal","Edge","Low",
    "Transaction FAILED/CANCELLED (not active).",
    "1. Record a manual Payment Entry.",
    "—",
    "Allowed (no active Pine Labs payment).")

# ── M. RECONCILE CRON ────────────────────────────────────────────────────
add("CRON-01","Reconcile","Terminal PENDING reconciled each minute","Happy","High",
    "A PENDING terminal txn with a Plutus id; scheduler enabled.",
    "1. Wait for the 1-minute cron.",
    "—",
    "Status pulled from the gateway; transaction finalized.")
add("CRON-02","Reconcile","Plural PENDING reconciled","Happy","High",
    "A PENDING Plural txn; scheduler enabled.",
    "1. Wait for the cron.",
    "—",
    "Status pulled from Plural; transaction finalized; realtime event refreshes the form.")
add("CRON-03","Reconcile","One bad row doesn't abort the batch","Edge","Medium",
    "Several PENDING rows, one causing an error.",
    "1. Run the reconcile task.",
    "—",
    "Per-row commit/rollback: good rows finalize; the bad row is logged; batch continues.")
add("CRON-04","Reconcile","Terminal row without payment_id skipped","Edge","Low",
    "A PENDING terminal row with no Plutus id.",
    "1. Run the reconcile task.",
    "—",
    "That row is skipped (nothing to query).")

# ── N. SECURITY ──────────────────────────────────────────────────────────
add("SEC-01","Security","Credentials stored encrypted","Security","High",
    "Machine/Plural credentials entered.",
    "1. Inspect Security Token / Plural Secret / Webhook Secret field types.",
    "—",
    "Stored as Password (encrypted at rest); not shown in plain text in the form.")
add("SEC-02","Security","Guest webhook endpoint is HMAC-gated","Security","High",
    "Webhook endpoint is allow_guest.",
    "1. POST without a valid signature.",
    "—",
    "Rejected (401); guest access is intentional but fail-closed on signature.")
add("SEC-03","Security","Secrets not leaked to logs by default","Security","Medium",
    "A terminal payment triggered; default log level.",
    "1. Trigger a payment.\n2. Grep logs for the Security Token.",
    "—",
    "Token not written at default (ERROR) log level. KNOWN ISSUE: BUG-09 — if INFO logging is enabled, the request payload (incl. token) is logged in cleartext.")

# ── O. AMOUNT / CONFIG EDGE ──────────────────────────────────────────────
add("AMT-01","Amount","Charge equals outstanding on submitted invoice","Boundary","High",
    "Submitted invoice with outstanding.",
    "1. Pay on Machine; observe the amount charged.",
    "outstanding=1000.",
    "Amount charged = outstanding_amount.")
add("AMT-02","Amount","Configured Amount Field is honored","Edge","Medium",
    "Payable row Amount Field set to a specific field.",
    "1. Set Amount Field to a non-default field with a distinct value.\n2. Pay.",
    "Amount Field = a custom field.",
    "Correct: charge should use the configured field. KNOWN ISSUE: BUG-11 — the field is ignored; outstanding/grand_total is used.")
add("AMT-03","Amount","Paise conversion rounds correctly","Boundary","Medium",
    "An amount whose paise conversion is float-sensitive.",
    "1. Pay ₹262.15 on the terminal.",
    "amount=262.15.",
    "Correct: 26215 paise captured. KNOWN ISSUE: BUG-06 — terminal truncates to 26214 (undercharge by 1 paisa).")
add("AMT-04","Amount","Zero/negative outstanding rejected","Negative","Medium",
    "An invoice with 0 outstanding.",
    "1. Attempt to initiate a payment.",
    "outstanding=0.",
    "Error: no outstanding amount; no transaction created.")

# ── Build workbook ───────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test Cases"
cols = ["TC ID","Module","Title","Type","Priority","Preconditions","Test Steps",
        "Test Data","Expected Result","Actual Result","Status"]
ws.append(cols)
for c in range(1, len(cols)+1):
    cell = ws.cell(1, c); cell.fill = HEAD_FILL; cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN
ws.row_dimensions[1].height = 26
ws.freeze_panes = "A2"

cur_mod = None
for row in T:
    _id, mod, title, typ, prio, pre, steps, data, exp = row
    if mod != cur_mod:
        # module banner row
        ws.append([mod] + [""]*(len(cols)-1))
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        bc = ws.cell(r, 1); bc.value = f"▮ {mod}"; bc.fill = MOD_FILL; bc.font = MOD_FONT
        bc.alignment = Alignment(vertical="center")
        cur_mod = mod
    ws.append([f"TC-{_id}", mod, title, typ, prio, pre, steps, data, exp, "", ""])
    r = ws.max_row
    for c in range(1, len(cols)+1):
        cell = ws.cell(r, c); cell.border = THIN
        cell.alignment = CENTER if c in (4,5) else WRAP
    tf = TYPE_FILL.get(typ)
    if tf: ws.cell(r,4).fill = PatternFill("solid", fgColor=tf)
    pf = PRIO_FILL.get(prio)
    if pf: ws.cell(r,5).fill = PatternFill("solid", fgColor=pf)
    # highlight known-issue expected cells
    if "KNOWN ISSUE" in exp:
        ws.cell(r,9).font = Font(color="9C0006")

widths = [12,16,34,12,10,34,46,20,50,26,12]
for i,w in enumerate(widths,1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── Summary sheet ────────────────────────────────────────────────────────
s = wb.create_sheet("Summary")
from collections import Counter, OrderedDict
by_mod = OrderedDict()
for row in T:
    by_mod.setdefault(row[1], Counter())[row[3]] += 1
types = ["Happy","Negative","Edge","Boundary","Security","Concurrency"]
s.append(["Module"] + types + ["Total"])
for c in range(1, len(types)+2):
    cell = s.cell(1,c); cell.fill=HEAD_FILL; cell.font=HEAD_FONT; cell.alignment=CENTER; cell.border=THIN
grand = Counter()
for mod, cnt in by_mod.items():
    tot = sum(cnt.values()); grand.update(cnt)
    s.append([mod] + [cnt.get(t,0) for t in types] + [tot])
    for c in range(1, len(types)+2):
        s.cell(s.max_row,c).border=THIN
        if c>1: s.cell(s.max_row,c).alignment=CENTER
s.append(["TOTAL"] + [grand.get(t,0) for t in types] + [sum(grand.values())])
for c in range(1, len(types)+2):
    cell=s.cell(s.max_row,c); cell.font=Font(bold=True); cell.border=THIN
    if c>1: cell.alignment=CENTER
s.column_dimensions["A"].width = 18
for i in range(2, len(types)+3):
    s.column_dimensions[get_column_letter(i)].width = 12

# ── Legend sheet ─────────────────────────────────────────────────────────
lg = wb.create_sheet("Legend")
lg.append(["How to use this workbook"]); lg["A1"].font = Font(bold=True, size=13)
notes = [
    "",
    "Test Cases sheet: one row per test case, grouped by Module (shaded banners).",
    "Execute each case, then fill 'Actual Result' and set 'Status' to Pass / Fail / Blocked.",
    "",
    "Type colours:",
    "  Happy       — main success path (green)",
    "  Negative    — invalid input / must be rejected (red)",
    "  Edge        — unusual-but-valid / recovery paths (yellow)",
    "  Boundary    — limits & value edges (blue)",
    "  Security    — auth / secrets / access (purple)",
    "  Concurrency — races / double-submit (teal)",
    "",
    "Priority: High / Medium / Low (business risk).",
    "",
    "Expected Result = the CORRECT/intended behavior.",
    "Cells marked 'KNOWN ISSUE: BUG-xx' (red text) describe where the app's CURRENT",
    "behavior deviates — those cases are expected to FAIL until the bug is fixed.",
    "Cross-reference pinelabs_bug_report.xlsx for the bug details.",
    "",
    "Environment: run against a site with pinelabs_fin + erpnext; a Pine Labs terminal",
    "simulator or sandbox for Terminal cases, and Plural sandbox creds for Payment Link cases.",
]
for n in notes:
    lg.append([n])
lg.column_dimensions["A"].width = 100

wb.save(OUT)
print(f"wrote {OUT}: {len(T)} test cases across {len(by_mod)} modules")
for mod,cnt in by_mod.items():
    print(f"  {mod:16} {sum(cnt.values()):3}")
