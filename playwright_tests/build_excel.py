"""Build pinelabs_bug_report.xlsx: Bugs sheet + Test Cases sheet.

Actual-result columns are left blank here; they are filled in by
record_results.py AFTER the Playwright run, so the workbook reflects
observed app behavior, not predictions.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = "pinelabs_bug_report.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
SEV1 = PatternFill("solid", fgColor="C00000")   # critical
SEV2 = PatternFill("solid", fgColor="ED7D31")   # high
SEV3 = PatternFill("solid", fgColor="FFC000")   # medium
SEV4 = PatternFill("solid", fgColor="A9D08E")   # low / feature
REFUTED = PatternFill("solid", fgColor="D9D9D9")
WHITE = Font(color="FFFFFF", bold=True)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"


# ── Bugs sheet ────────────────────────────────────────────────────────────
# id, title, severity(1-4), area, file:line, mechanism, impact, ui_testable, status
BUGS = [
    ("BUG-01", "Malformed HTTP-200 body reported as success", 1, "Terminal client",
     "api/pinelabs_client.py:265",
     "_parse_response returns success=True, data={} when the 200 body is not JSON; terminal.py then falls into the 'Synchronous decline' branch (response_code=None).",
     "Terminal is live and customer pays; ERPNext marks the txn FAILED with no PlutusReference to recover it. Unreconcilable.",
     "No (needs gateway to return non-JSON; backend harness)", "Confirmed (backend)"),
    ("BUG-02", "Plural uses invoice name as merchant reference", 1, "Payment link",
     "api/plural.py:250",
     "merchant_payment_link_reference = source_doc.name (invoice), not txn name. A 2nd link on the same invoice shares the reference; status lookup can't tell attempts apart.",
     "Cron reads the old attempt's EXPIRED status, marks the new txn FAILED; customer pays the live link, webhook mark_success throws on FAILED. Charged, invoice unpaid.",
     "No (needs Plural gateway)", "Reported (not reproduced)"),
    ("BUG-03", "PENDING->FAILED/CANCELLED not guarded -> Retry double-charges", 1, "State machine",
     "doctype/pinelabs_transaction/pinelabs_transaction.py:177",
     "Controller flag-gates only SUCCESS. FAILED/CANCELLED from PENDING are allowed with no internal flag, and status is an editable Select with System Manager write.",
     "Manual FAILED on a live PENDING txn -> mark_success throws, no PE; cron skips (PENDING-only); dialog offers Retry -> second charge.",
     "Yes (Desk form)", "Confirmed (backend)"),
    ("BUG-04", "Retry drops create_payment_entry override", 2, "Terminal flow",
     "api/terminal.py:424",
     "retry_failed_terminal_payment re-initiates with only doctype/name/mode; the stored create_payment_entry_override is not forwarded.",
     "Custom doctype paid via Retry gets no Payment Entry; customer charged, nothing booked. Pre-flight checks skipped too.",
     "No (needs API caller w/ override)", "Reported (not reproduced)"),
    ("BUG-05", "cancel_raw omits ClientID/StoreID", 1, "Terminal client",
     "api/pinelabs_client.py:428",
     "Cancel/force-cancel payloads omit ClientID and StoreID that every sibling call sends and the constructor treats as mandatory. NEEDS confirmation vs Pine Labs spec.",
     "If gateway rejects the cancel: txn marked CANCELLED while terminal stays armed; late tap charges the customer with no reconciling row.",
     "No (needs gateway)", "Uncertain - verify vs spec"),
    ("BUG-06", "Amount truncated (int) not rounded -> undercharge by 1 paisa", 2, "Terminal client",
     "api/pinelabs_client.py:335",
     "amount_in_paisa = int(float(amount)*100) truncates binary-float error. 4.59% of 2-dp values lose a paisa. plural.py:301 already uses round().",
     "Terminal captures a paisa less than the booked Payment Entry (rupee figure). Invoice reads fully paid; settlement permanently short.",
     "No (backend math)", "Confirmed (backend)"),
    ("BUG-07", "Same as BUG-04 (retry/override) - merged", 2, "Terminal flow",
     "api/terminal.py:424", "Duplicate of BUG-04.", "See BUG-04.",
     "No", "Merged into BUG-04"),
    ("BUG-08", "No double-submit guard on Start/Retry", 2, "Buttons JS / server",
     "public/js/pinelabs_buttons.js:284; api/transaction.py:176",
     "Neither dialog button disables on click; server guard is check-then-act SELECT with no lock/unique constraint.",
     "Double-click queues two UploadBilledTransaction requests to the same terminal; relies on terminal to reject the 2nd.",
     "Partly (double-click timing, flaky)", "Reported (not reproduced)"),
    ("BUG-09", "SecurityToken logged in cleartext", 3, "Terminal client",
     "api/pinelabs_client.py:308,282",
     "Full request+response payloads logged at INFO; payload carries SecurityToken (a Password field) with MerchantID/ClientID/StoreID.",
     "Anyone with log access or a log pipeline harvests the full terminal credential set.",
     "No (log inspection)", "Confirmed (code)"),
    ("BUG-10", "Webhook-secret length check reads old secret", 3, "Settings",
     "doctype/pinelabs_settings/pinelabs_settings.py:76",
     "Claimed get_password() during validate() reads the previously-stored secret, bypassing the >=16 rule.",
     "Claimed: a weak webhook HMAC secret could be saved.",
     "Yes (Settings form)", "REFUTED (live test)"),
    ("BUG-11", "Configured Amount Field is never read", 4, "Config / amount",
     "api/transaction.py:913 (loaded, unused)",
     "amount_field is mandatory + has a UI picker but no resolver consumes it; all paths hardcode outstanding_amount/grand_total/amount.",
     "A custom doctype whose amount is elsewhere never charges (button gate computes 0; API throws 'no outstanding').",
     "Hard (needs custom doctype)", "Confirmed (code)"),
    ("BUG-12", "Waiting-modal registry drops caller callbacks", 4, "Buttons JS",
     "public/js/pinelabs_buttons.js:344",
     "Registry hit returns the existing modal without merging on_settled/on_cancel/poll; the global realtime listener opens it option-less first.",
     "on_settled never fires (invoice doesn't refresh; POS hangs); external-poll flag lost -> double polling.",
     "Hard (realtime race)", "Reported (not reproduced)"),
    ("BUG-13", "Buttons never render on drafts", 4, "Buttons JS",
     "public/js/pinelabs_buttons.js:246",
     "Claimed outstanding_amount=0 on drafts makes the amount gate bail. Depends on ERPNext populating outstanding_amount on a saved draft.",
     "Claimed: draft path (show_on_docstatus incl. 0) unreachable from the UI.",
     "Yes (draft SI form)", "Likely REFUTED (verify)"),
    ("BUG-14", "Dynamic Link resolves against hardcoded doctype", 4, "Contact resolver",
     "api/customer_contact.py:163",
     "Linked-DocType lookup uses a static configured doctype, ignoring the doc's real Dynamic Link target.",
     "party_type=Supplier with a Customer-pinned map returns a different party's mobile; pay-link SMS'd to wrong person.",
     "No (backend resolver)", "Confirmed (backend)"),
    ("BUG-15", "Failed webhook swallowed by dedup row", 4, "Webhook",
     "api/webhook.py:74,135",
     "Event row inserted 'received' before processing; a processing exception returns 500 (committed), so the retry hits the unique constraint and 200s without processing.",
     "Primary webhook signal silently lost; relies on the 1-min cron fallback; event stuck 'received'.",
     "No (needs webhook delivery)", "Confirmed (code)"),
]

# ── Test cases (human-usable UI flows + the backend-only notes) ─────────────
# id, bug, title, precondition, steps, expected(correct behavior), type
TCS = [
    ("TC-01", "BUG-13", "Buttons on a submitted UNPAID invoice (baseline)",
     "Sales Invoice ACC-SINV-2026-00003 submitted, outstanding 1000; SI payable row enabled.",
     "1. Log in to Desk. 2. Open /app/sales-invoice/ACC-SINV-2026-00003. 3. Wait for form load. 4. Look for 'Pay on Machine' and 'Send Payment Link' buttons in the toolbar.",
     "Both buttons appear (positive control).", "UI / Playwright"),
    ("TC-02", "BUG-13", "Buttons on a PAID invoice should be hidden",
     "Sales Invoice ACC-SINV-2026-00001 is Paid (outstanding 0).",
     "1. Open /app/sales-invoice/ACC-SINV-2026-00001. 2. Wait for load. 3. Look for the two buttons.",
     "Neither button appears (amount gate correctly hides them when outstanding=0).", "UI / Playwright"),
    ("TC-03", "BUG-13", "Buttons on a DRAFT invoice (the bug claim)",
     "Sales Invoice ACC-SINV-2026-00002 is Draft; SI row show_on_docstatus='0,1' so drafts are allowed.",
     "1. Open /app/sales-invoice/ACC-SINV-2026-00002. 2. Confirm status = Draft (docstatus 0). 3. Look for the two buttons.",
     "Per BUG-13: buttons should be ABSENT on a draft. (Observe whether they actually appear.)", "UI / Playwright"),
    ("TC-04", "BUG-03", "Manually set a PENDING transaction to FAILED",
     "Pinelabs Transaction PLT-2026-00002 exists in PENDING (Terminal).",
     "1. Open /app/pinelabs-transaction/PLT-2026-00002. 2. Change Status field from PENDING to FAILED. 3. Save (Ctrl+S). 4. Observe accept/reject.",
     "Correct behavior: the save is REJECTED (only the txn service may drive transitions). Observe whether it is actually accepted.", "UI / Playwright"),
    ("TC-05", "BUG-03", "Manually set a transaction to SUCCESS (control)",
     "A PENDING/INITIATED Pinelabs Transaction open in the form.",
     "1. Open the transaction. 2. Change Status to SUCCESS. 3. Save. 4. Observe.",
     "Correct behavior: REJECTED with 'Status SUCCESS may only be set by the Pinelabs transaction service'. (This guard is expected to work.)", "UI / Playwright"),
    ("TC-06", "BUG-03", "Pay on Machine pre-flight error (no account mapped)",
     "PineLabs modes have no default account for Pine Lab Machine; submitted unpaid SI.",
     "1. Open ACC-SINV-2026-00003. 2. Click 'Pay on Machine'. 3. In the dialog pick 'PineLabs - Card'. 4. Click Start Payment. 5. Observe the message.",
     "Fail-fast: a clear 'no default account configured' error appears BEFORE any gateway call (positive - protects the customer).", "UI / Playwright"),
    ("TC-07", "BUG-06", "Amount to paise conversion",
     "N/A - server-side arithmetic.",
     "Call the exact expression int(float(amount)*100) for representative amounts (262.15, 8.20, 1.15).",
     "Correct: 26215, 820, 115. Bug: 26214, 819, 114.", "Backend harness"),
    ("TC-08", "BUG-01", "Non-JSON 200 response classification",
     "N/A - requires the gateway to return a non-JSON 200.",
     "Feed a 200/HTML response to PineLabsClient._parse_response and trace terminal.initiate branching.",
     "Correct: treated as failure. Bug: success=True -> 'Synchronous decline' FAILED.", "Backend harness"),
    ("TC-09", "BUG-14", "Dynamic Link contact resolution",
     "Customer & Supplier share a name with different mobiles; map pinned to Customer.",
     "Call customer_contact._resolve_value on a doc with party_type=Supplier.",
     "Correct: resolve the Supplier (or refuse). Bug: returns the Customer's mobile.", "Backend harness"),
]


def build():
    wb = openpyxl.Workbook()

    # Bugs
    ws = wb.active
    ws.title = "Bugs"
    cols = ["Bug ID", "Title", "Severity", "Area", "Location", "Mechanism",
            "Impact", "UI-testable?", "Status"]
    ws.append(cols)
    sevfill = {1: SEV1, 2: SEV2, 3: SEV3, 4: SEV4}
    for b in BUGS:
        bid, title, sev, area, loc, mech, impact, uit, status = b
        ws.append([bid, title, f"S{sev}", area, loc, mech, impact, uit, status])
        r = ws.max_row
        cell = ws.cell(row=r, column=3)
        cell.fill = sevfill[sev]
        cell.font = WHITE
        cell.alignment = Alignment(horizontal="center", vertical="top")
        if "REFUTED" in status:
            for c in range(1, len(cols) + 1):
                ws.cell(row=r, column=c).fill = REFUTED
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            cell.alignment = WRAP if c in (6, 7) else TOP
    widths = [9, 34, 9, 15, 34, 52, 46, 26, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    style_header(ws, len(cols))

    # Test Cases (Actual + Verdict blank until Playwright run)
    ws2 = wb.create_sheet("Test Cases")
    cols2 = ["TC ID", "Bug", "Title", "Precondition", "Steps",
             "Expected Result", "Type", "Actual Result (observed)", "Verdict"]
    ws2.append(cols2)
    for t in TCS:
        ws2.append([*t, "", ""])   # actual + verdict filled post-run
        r = ws2.max_row
        for c in range(1, len(cols2) + 1):
            cell = ws2.cell(row=r, column=c)
            cell.border = THIN
            cell.alignment = WRAP if c in (4, 5, 6, 8) else TOP
    widths2 = [8, 9, 34, 34, 52, 44, 16, 50, 26]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    style_header(ws2, len(cols2))

    wb.save(OUT)
    print(f"wrote {OUT}: Bugs={len(BUGS)} rows, Test Cases={len(TCS)} rows")


if __name__ == "__main__":
    build()
